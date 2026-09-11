import openpyxl
from datetime import datetime
from openpyxl.utils.datetime import to_excel
src=openpyxl.load_workbook('EXCEL UTILITARIOS/Formato Libro Recaudaciones.xlsx')
path='outputs/recaudaciones/Formato Libro Recaudaciones.xlsx'
out=openpyxl.load_workbook(path,data_only=True)
forms=openpyxl.load_workbook(path,data_only=False)
dh={c.value.lower():c.column for c in src['Datos'][1]}
errors=[]
for row in src['Datos']:
 for c in row:
  a,b=c.value,out['Datos'][c.coordinate].value
  if a in (None,'') and b in (None,''):continue
  if a!=b: errors.append(('source',c.coordinate,str(a),str(b)))
for j in range(1,87):
 h=src['Formato'].cell(1,j).value
 assert forms['Formato'].cell(1,j).value==h
 for i in range(2,15):
  a=src['Datos'].cell(i,dh[h.lower()]).value
  b=out['Formato'].cell(i,j).value
  if isinstance(a,datetime): a=to_excel(a)
  if isinstance(b,datetime): b=to_excel(b)
  if a in (None,'') and b in (None,''):continue
  if isinstance(a,(int,float)) and isinstance(b,(int,float)) and abs(a-b)<1e-8:continue
  if a!=b:errors.append(('mapped',i,j,str(a),str(b)))
print('Errors',len(errors),errors[:8],errors[-12:])
print('Tables',[(t.name,t.ref) for t in forms['Formato'].tables.values()])
print('Formula count',sum(c.data_type=='f' for row in forms['Formato'].iter_rows(max_col=86) for c in row))
assert not errors

import zipfile, os
from lxml import etree as E
from datetime import datetime
import openpyxl
from openpyxl.utils.datetime import to_excel
from openpyxl.utils import get_column_letter
p='outputs/recaudaciones/Formato Libro Recaudaciones.xlsx'
src=openpyxl.load_workbook('EXCEL UTILITARIOS/Formato Libro Recaudaciones.xlsx')
ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
q=lambda x:'{'+ns+'}'+x
def put(c,value,formula=None):
 for child in list(c):
  if child.tag in (q('v'),q('is'),q('f')):c.remove(child)
 c.attrib.pop('t',None)
 if formula is not None:E.SubElement(c,q('f')).text=formula
 if isinstance(value,datetime):value=to_excel(value)
 if value is None or value=='':
  if formula is not None:c.set('t','str');E.SubElement(c,q('v')).text=''
 elif isinstance(value,bool):c.set('t','b');E.SubElement(c,q('v')).text='1' if value else '0'
 elif isinstance(value,(float,int)):E.SubElement(c,q('v')).text=str(value)
 elif formula is not None:c.set('t','str');E.SubElement(c,q('v')).text=value
 else:
  c.set('t','inlineStr');E.SubElement(E.SubElement(c,q('is')),q('t')).text=value
with zipfile.ZipFile(p) as z:
 payload={n:z.read(n) for n in z.namelist()}
for idx,name in [(1,'Datos'),(2,'Formato')]:
 key=f'xl/worksheets/sheet{idx}.xml'; root=E.fromstring(payload[key]); cells={c.get('r'):c for c in root.iter(q('c'))}
 if name=='Datos':
  for row in src[name]:
   for c in row:
    if c.coordinate in cells:put(cells[c.coordinate],c.value)
 else:
  headers={c.value.lower():c.column for c in src['Datos'][1]}
  for j in range(1,87):
   k=headers[src[name].cell(1,j).value.lower()]
   for i in range(2,15):
    ref=f"'Datos'!{get_column_letter(k)}{i}"
    put(cells[f'{get_column_letter(j)}{i}'],src['Datos'].cell(i,k).value,f'IF(ISBLANK({ref}),"",{ref})')
 payload[key]=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(p+'.tmp','w',zipfile.ZIP_DEFLATED) as z:
 for n,data in payload.items():z.writestr(n,data)
os.replace(p+'.tmp',p)
print('Repaired blank and boolean serialization; original values and formulas retained.')
